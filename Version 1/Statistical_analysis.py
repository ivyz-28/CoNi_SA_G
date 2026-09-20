from ase.io.trajectory import Trajectory
from pathlib import Path
import os
import pandas as pd
import openpyxl
from openpyxl import Workbook
import numpy as np
import re
import matplotlib.pyplot as plt
from sklearn.mixture import GaussianMixture

metals = [] # List all metals to carry out statistical analysis for 
out_direc = "" # Replace with desired output directory
os.makedirs(f'{out_direct}/Fairchem Excel', exist_ok=True)
os.makedirs(f'{out_direct}/Fairchem Plots', exist_ok=True)

def split_label(s):
    match = re.match(r"([a-zA-Z]+)(\d+)", s)
    if not match:
        raise ValueError(f"'{s}' doesn't match the expected letters+numbers pattern")
    return match.group(1), match.group(2)

df1 = pd.read_excel(f'',usecols = ['E_catalyst','Energy'])
df1.itertuples()
sa = np.array([x for x in df1.E_catalyst if not pd.isnull(x)])
energy = np.array([x for x in df1.Energy if not pd.isnull(x)])
single_atom_dict = {}

for i, atom in enumerate(sa):
    single_atom_dict[atom.split('.traj')[0]] = energy[i]

n_components = 3
colors = plt.cm.tab20.colors
gmm_store={}

for metal in metals:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{metal} Lattice Fairchem Energies"

    ws['A1'] = "Structure_with_H"
    ws['B1'] = "E_H"
    ws['C1'] = "Structure_without_H"
    ws['D1'] = "E_catalyst"
    ws['E1'] = "n_Co"
    ws['F1'] = "n_Ni"
    ws['G1'] = "n_Pt"
    ws['H1'] = "E_coh"
    ws['I1'] = "Ground_state"
    ws['J1'] = "p_i"
    ws['K1'] = "E_ads"
    ws['L1'] = "E_ads_final"
    

    folder_path = f'./{metal} Fairchem'

    h_dict = {}
    no_h_dict = {}

    

    for doc in os.listdir(folder_path):
        if doc.endswith('traj'):
            print(doc)
            traj = Trajectory(f'{folder_path}/{doc}')
            last_frame = traj[-1]
            energy = last_frame.get_potential_energy()
            if "H" in doc:
                h_dict[doc] = energy
            else:
                no_h_dict[doc] = energy

    sorted_h = dict(sorted(h_dict.items()))
    sorted_no_h = dict(sorted(no_h_dict.items()))
    energy_coh_list = []

    for i, key in enumerate(sorted_h):
        ws.cell(row = i+2, column = 1, value = key)
        ws.cell(row = i+2, column = 2, value = sorted_h[key])
    for i, key in enumerate(sorted_no_h):
        ws.cell(row = i+2, column = 3, value = key)
        ws.cell(row = i+2, column = 4, value = sorted_no_h[key])
        file = f"{folder_path}/{key.split('.traj')[0]}.vasp"
        with open(file, 'r') as f:
            number_dict = {}
            for atom in f.readlines()[0].split():
                symbols, numbers = split_label(atom)
                number_dict[symbols] = int(numbers)
        ws.cell(row = i+2, column = 5, value = number_dict["Co"])
        ws.cell(row = i+2, column = 6, value = number_dict["Ni"])
        ws.cell(row = i+2, column = 7, value = number_dict[metal])

        
        total_num = number_dict["Co"] + number_dict["Ni"] + number_dict[metal]
        energy_coh = (sorted_no_h[key] - number_dict["Co"] * single_atom_dict["Co"] - number_dict["Ni"] * single_atom_dict["Ni"] - number_dict[metal] * single_atom_dict[metal]) / total_num
        ws.cell(row = i+2, column = 8, value = energy_coh)
        energy_coh_list.append(energy_coh)
    print(energy_coh_list)
    base_energy = min(energy_coh_list)
    ws.cell(row = 2, column = 9, value = base_energy)

    wb.save(filename = f"{out_direct}/Fairchem Excel/{metal} Fairchem Lattice Energies.xlsx")
    wb.close()

    df = pd.read_excel(f"{out_direct}/Fairchem Excel/{metal} Fairchem Lattice Energies.xlsx", usecols=['E_H', 'E_catalyst', 'E_coh'])
    df.itertuples()
    energy_H = np.array([x for x in df.E_H if not pd.isnull(x)])
    energy_no_H = np.array([x for x in df.E_catalyst if not pd.isnull(x)])
    cohesive_energies = np.array([x for x in df.E_coh if not pd.isnull(x)])

    exponential = -1*np.exp((base_energy - cohesive_energies)/(3200*0.00008617))
    partition = sum(exponential)
    probabilities = exponential / partition

    wb = openpyxl.load_workbook(f"{out_direct}/Fairchem Excel/{metal} Fairchem Lattice Energies.xlsx")
    ws = wb.active

    for i, probability in enumerate(probabilities):
        ws.cell(row = i+2, column = 10, value = probability)
        
    E_ads = energy_H - energy_no_H - 0.5*(-6.940629583)
    for i, value in enumerate(E_ads):
        ws.cell(row = i+2, column = 11, value = value)
    ads_final = sum(E_ads * probabilities)
    ws.cell(row = 2, column = 12, value = ads_final)
    wb.save(filename = f"{out_direct}/Fairchem Excel/{metal} Fairchem Lattice Energies.xlsx")
    wb.close()

    idx = np.random.choice(
    len(E_ads),
    size=10000,
    p=probabilities
    )
    
    gmm = GaussianMixture(n_components=n_components, random_state=0)
    gmm.fit(E_ads[idx].reshape(-1, 1))
    

    gmm_store[metal] = {'gmm': gmm, 'k': ads_final,
                      'E_range': (E_ads.min(), E_ads.max())}

    ads_range = np.max(E_ads) - np.min(E_ads)
    fig,ax = plt.subplots(figsize = (10,5))

    n, bins, patches = ax.hist(
        E_ads,
        bins=int(ads_range/0.025),
        weights=probabilities,
        color='steelblue',
        edgecolor='grey',
        linewidth=0.4,
    )
    for patch, left_edge in zip(patches, bins[:-1]):
        if -0.35 <= left_edge < -0.05:
            patch.set_facecolor('red')
    

    x_grid    = np.linspace(bins[0], bins[-1], 1000).reshape(-1, 1)
    bin_width = bins[1] - bins[0]
    pdf       = np.exp(gmm.score_samples(x_grid)) * bin_width
    x_range  = np.linspace(-0.35, -0.05, 1000).reshape(-1, 1)
    pdf_area = np.exp(gmm.score_samples(x_range))
    area     = np.trapezoid(pdf_area.flatten(), x_range.flatten())

    ax.plot(x_grid, pdf, color='black', linewidth=2,
            label=f'GMM (n={n_components})  |  P(−0.35→−0.05) = {area:.3f}')
        
    ax.axvline(x=0, color='black', linewidth=0.8, linestyle='-')        
    ax.axvline(x=ads_final, color='darkblue', linewidth=1.5, linestyle='-', label=f'E_ads = {ads_final:.2f}')
    ax.legend(fontsize=11)

    ax.set_xlabel('E_ads (eV)', fontsize=12)
    ax.set_ylabel('Probability', fontsize=12)
    ax.set_title(f'{metal} Catalyst — Adsorption Energy Distribution (Fairchem)', fontsize=13)
    plt.tight_layout()
    plt.savefig(f'{out_direct}/Fairchem Plots/{metal}.pdf', dpi=300, bbox_inches='tight')
    plt.savefig(f'{out_direct}/Fairchem Plots/{metal}.png', dpi=300, bbox_inches='tight')
    plt.close()
                
